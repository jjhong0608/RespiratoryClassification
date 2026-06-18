from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from collections.abc import Mapping, Sequence
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.utils.fs import Fs
from src.utils.logging import LoggingMixin, enable_file_logging, logger

DEFAULT_DATA_ROOT = Path(
    "/Users/jjhong0608/Documents/AudioData/RespiratoryClassification/DATA/"
    "DISEASE_CNUH_DATA"
)
DEFAULT_EXCEL_ROOT = Path(
    "/Users/jjhong0608/Documents/AudioData/RespiratoryClassification/DATA/EXCEL/"
    "Disease_Data"
)
DEFAULT_OUT_DIR = DEFAULT_DATA_ROOT / "disease_dataset_catalog"

PRIMARY_LABELS = ("Normal", "Airway", "Lung_Parenchymal")
SPLITS = ("fold_0", "fold_1", "fold_2", "fold_3", "fold_4", "test")

AUSCULTATION_COLUMN = (
    "청진음 \n"
    "1: Crackle, 2: wheezing, \n"
    "3: rhonchi, 4: stridor, \n"
    "5: non specific, 6:crackle+wheezing \n"
    "7:crackle+rhonchi, \n"
    "8:wheezing + rhonchi, \n"
    "9:crackle+wheezing,+rhonchi\n"
    "10:wheezing +stridor 11: ETC \n"
    "12: decreased BS\n"
    "99. Obscure"
)
DIAGNOSIS_COLUMN = (
    "1: pneumonia, \n"
    "2: IPF, 3: COPD, \n"
    "4: asthma, \n"
    "5: lung cancer, \n"
    "6: healthy 7: TB, \n"
    "8: Bronchiectasis, \n"
    "9; asthma +pneumonia, \n"
    "10: COPD+pneumonia, \n"
    "11: ILD except IPF (CPFE포함) \n"
    "12: COPD+ IPF, 13. ACO \n"
    "14: ETC, 15. lung nodule \n"
    "99: no diagnosis)"
)

AUSCULTATION_LABELS = {
    1: "Crackle",
    2: "wheezing",
    3: "rhonchi",
    4: "stridor",
    5: "non specific",
    6: "crackle+wheezing",
    7: "crackle+rhonchi",
    8: "wheezing+rhonchi",
    9: "crackle+wheezing+rhonchi",
    10: "wheezing+stridor",
    11: "ETC",
    12: "decreased BS",
    99: "Obscure",
}
DIAGNOSIS_LABELS = {
    1: "pneumonia",
    2: "IPF",
    3: "COPD",
    4: "asthma",
    5: "lung cancer",
    6: "healthy",
    7: "TB",
    8: "Bronchiectasis",
    9: "asthma+pneumonia",
    10: "COPD+pneumonia",
    11: "ILD except IPF (CPFE included)",
    12: "COPD+IPF",
    13: "ACO",
    14: "ETC",
    15: "lung nodule",
    99: "no diagnosis",
}

CATALOG_COLUMNS = [
    "filename_stem",
    "filename",
    "audio_path",
    "primary_label",
    "binary_label",
    "is_abnormal",
    "excel_source_file",
    "excel_row_number",
    "auscultation_raw",
    "auscultation_code",
    "auscultation_label",
    "diagnosis_raw",
    "diagnosis_code",
    "diagnosis_label",
    "split",
    "split_role",
    "cv_validation_fold",
    "is_test",
    "split_primary_label",
    "split_audio_path",
    "appears_in_abnormal_overlay",
    "abnormal_overlay_path",
    "has_excel_metadata",
    "has_split_membership",
    "label_matches_excel_source",
    "label_matches_split_folder",
    "notes",
]

WARNING_COLUMNS = [
    "warning_type",
    "filename_stem",
    "filename",
    "primary_label",
    "source_label",
    "path",
    "excel_source_file",
    "excel_row_number",
    "column",
    "raw_value",
    "message",
]


@dataclass(frozen=True)
class WavRecord:
    filename_stem: str
    filename: str
    audio_path: Path
    primary_label: str


@dataclass(frozen=True)
class ParsedAnnotation:
    raw: str
    code: int | None
    label: str


@dataclass(frozen=True)
class ExcelAnnotationRecord:
    filename_stem: str
    filename: str
    source_label: str
    excel_source_file: Path
    excel_row_number: int
    auscultation: ParsedAnnotation
    diagnosis: ParsedAnnotation


@dataclass(frozen=True)
class SplitMembership:
    split: str
    split_role: str
    cv_validation_fold: int | None
    is_test: bool
    split_primary_label: str
    split_audio_path: Path
    appears_in_abnormal_overlay: bool
    abnormal_overlay_path: Path | None


@dataclass(frozen=True)
class CatalogWarning:
    warning_type: str
    filename_stem: str = ""
    filename: str = ""
    primary_label: str = ""
    source_label: str = ""
    path: str = ""
    excel_source_file: str = ""
    excel_row_number: int | str = ""
    column: str = ""
    raw_value: str = ""
    message: str = ""


@dataclass(frozen=True)
class CatalogBuildResult:
    rows: list[dict[str, Any]]
    warnings: list[CatalogWarning]
    summary: dict[str, Any]


class AnnotationCodeParser:
    CODE_PATTERN = re.compile(r"\d+")

    def __init__(self, *, code_map: Mapping[int, str], column_name: str):
        self.code_map = dict(code_map)
        self.column_name = column_name

    def parse(self, value: object) -> ParsedAnnotation:
        raw = "" if value is None else str(value).strip()
        match = self.CODE_PATTERN.search(raw)
        if match is None:
            return ParsedAnnotation(raw=raw, code=None, label="")
        code = int(match.group(0))
        return ParsedAnnotation(raw=raw, code=code, label=self.code_map.get(code, ""))


class WavInventoryLoader(LoggingMixin):
    def __init__(self, data_root: Path):
        self.data_root = data_root

    def load(self) -> dict[str, WavRecord]:
        records: dict[str, WavRecord] = {}
        seen_labels: defaultdict[str, list[str]] = defaultdict(list)
        for label in PRIMARY_LABELS:
            label_dir = self.data_root / label
            if not label_dir.is_dir():
                raise FileNotFoundError(
                    f"Primary label directory not found: {label_dir}"
                )
            for audio_path in sorted(label_dir.glob("*.wav")):
                stem = audio_path.stem
                seen_labels[stem].append(label)
                records[stem] = WavRecord(
                    filename_stem=stem,
                    filename=audio_path.name,
                    audio_path=audio_path,
                    primary_label=label,
                )
        duplicates = {
            stem: labels for stem, labels in seen_labels.items() if len(labels) > 1
        }
        if duplicates:
            examples = dict(list(sorted(duplicates.items()))[:10])
            raise ValueError(
                "Primary wav filename appears under more than one primary label: "
                f"{examples}"
            )
        self.logger.info("Loaded %d primary wav records", len(records))
        return records


class ExcelAnnotationLoader(LoggingMixin):
    def __init__(self, excel_root: Path):
        self.excel_root = excel_root
        self.auscultation_parser = AnnotationCodeParser(
            code_map=AUSCULTATION_LABELS,
            column_name=AUSCULTATION_COLUMN,
        )
        self.diagnosis_parser = AnnotationCodeParser(
            code_map=DIAGNOSIS_LABELS,
            column_name=DIAGNOSIS_COLUMN,
        )

    def load(self) -> tuple[dict[str, ExcelAnnotationRecord], list[CatalogWarning]]:
        annotations: dict[str, ExcelAnnotationRecord] = {}
        warnings: list[CatalogWarning] = []
        for label in PRIMARY_LABELS:
            path = self.excel_root / f"{label}.xlsx"
            if not path.is_file():
                raise FileNotFoundError(f"Required Excel file not found: {path}")
            label_records, label_warnings = self._load_label_file(label, path)
            warnings.extend(label_warnings)
            for record in label_records:
                if record.filename_stem in annotations:
                    previous = annotations[record.filename_stem]
                    raise ValueError(
                        "Excel filename appears in multiple rows/files: "
                        f"{record.filename_stem} | {previous.excel_source_file}:"
                        f"{previous.excel_row_number} and {path}:{record.excel_row_number}"
                    )
                annotations[record.filename_stem] = record
        self.logger.info("Loaded %d Excel annotation records", len(annotations))
        return annotations, warnings

    def _load_label_file(
        self,
        label: str,
        path: Path,
    ) -> tuple[list[ExcelAnnotationRecord], list[CatalogWarning]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook[workbook.sheetnames[0]]
            header_row = next(
                worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
            )
            headers = {
                value: index
                for index, value in enumerate(header_row)
                if value is not None
            }
            self._require_columns(headers, path)
            filename_index = headers["filename"]
            auscultation_index = headers[AUSCULTATION_COLUMN]
            diagnosis_index = headers[DIAGNOSIS_COLUMN]
            records: list[ExcelAnnotationRecord] = []
            warnings: list[CatalogWarning] = []
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True),
                start=2,
            ):
                filename_raw = row[filename_index]
                if filename_raw is None or str(filename_raw).strip() == "":
                    continue
                stem = normalize_filename_stem(filename_raw)
                auscultation = self.auscultation_parser.parse(row[auscultation_index])
                diagnosis = self.diagnosis_parser.parse(row[diagnosis_index])
                records.append(
                    ExcelAnnotationRecord(
                        filename_stem=stem,
                        filename=f"{stem}.wav",
                        source_label=label,
                        excel_source_file=path,
                        excel_row_number=row_number,
                        auscultation=auscultation,
                        diagnosis=diagnosis,
                    )
                )
                warnings.extend(
                    self._annotation_warnings(
                        label=label,
                        path=path,
                        row_number=row_number,
                        stem=stem,
                        auscultation=auscultation,
                        diagnosis=diagnosis,
                    )
                )
            return records, warnings
        finally:
            workbook.close()

    @staticmethod
    def _require_columns(headers: Mapping[object, int], path: Path) -> None:
        missing = [
            column
            for column in ("filename", AUSCULTATION_COLUMN, DIAGNOSIS_COLUMN)
            if column not in headers
        ]
        if missing:
            raise ValueError(f"Required Excel column(s) missing in {path}: {missing}")

    @staticmethod
    def _annotation_warnings(
        *,
        label: str,
        path: Path,
        row_number: int,
        stem: str,
        auscultation: ParsedAnnotation,
        diagnosis: ParsedAnnotation,
    ) -> list[CatalogWarning]:
        warnings: list[CatalogWarning] = []
        for column, parsed in (
            ("auscultation", auscultation),
            ("diagnosis", diagnosis),
        ):
            if parsed.code is None:
                warnings.append(
                    CatalogWarning(
                        warning_type="annotation_code_unparsed",
                        filename_stem=stem,
                        filename=f"{stem}.wav",
                        source_label=label,
                        excel_source_file=str(path),
                        excel_row_number=row_number,
                        column=column,
                        raw_value=parsed.raw,
                        message="Could not parse an integer annotation code",
                    )
                )
            elif not parsed.label:
                warnings.append(
                    CatalogWarning(
                        warning_type="annotation_code_unknown",
                        filename_stem=stem,
                        filename=f"{stem}.wav",
                        source_label=label,
                        excel_source_file=str(path),
                        excel_row_number=row_number,
                        column=column,
                        raw_value=parsed.raw,
                        message=f"Parsed code {parsed.code} is not in the configured code map",
                    )
                )
        return warnings


class SplitMembershipLoader(LoggingMixin):
    def __init__(self, data_root: Path):
        self.data_root = data_root
        self.folds_root = data_root / "5_Folds"

    def load(
        self,
        wav_inventory: Mapping[str, WavRecord],
    ) -> dict[str, SplitMembership]:
        memberships: dict[str, SplitMembership] = {}
        for split in SPLITS:
            split_dir = self.folds_root / split
            if not split_dir.is_dir():
                raise FileNotFoundError(
                    f"Required split directory not found: {split_dir}"
                )
            abnormal_index = {
                audio_path.stem: audio_path
                for audio_path in sorted((split_dir / "Abnormal").glob("*.wav"))
            }
            for label in PRIMARY_LABELS:
                label_dir = split_dir / label
                if not label_dir.is_dir():
                    raise FileNotFoundError(
                        f"Required split label directory not found: {label_dir}"
                    )
                for split_audio_path in sorted(label_dir.glob("*.wav")):
                    stem = split_audio_path.stem
                    if stem not in wav_inventory:
                        raise ValueError(
                            "Split primary file is not found in root primary wav inventory: "
                            f"{split_audio_path}"
                        )
                    if stem in memberships:
                        previous = memberships[stem]
                        raise ValueError(
                            "Primary wav appears in multiple split memberships: "
                            f"{stem} | {previous.split} and {split}"
                        )
                    abnormal_path = abnormal_index.get(stem)
                    memberships[stem] = SplitMembership(
                        split=split,
                        split_role="test" if split == "test" else "validation",
                        cv_validation_fold=None
                        if split == "test"
                        else int(split.rsplit("_", 1)[1]),
                        is_test=split == "test",
                        split_primary_label=label,
                        split_audio_path=split_audio_path,
                        appears_in_abnormal_overlay=abnormal_path is not None,
                        abnormal_overlay_path=abnormal_path,
                    )
        self.logger.info("Loaded %d split memberships", len(memberships))
        return memberships


class AbnormalOverlayValidator(LoggingMixin):
    def __init__(self, data_root: Path):
        self.data_root = data_root

    def validate(self) -> None:
        self._validate_dir(self.data_root)
        for split in SPLITS:
            self._validate_dir(self.data_root / "5_Folds" / split)

    def _validate_dir(self, root: Path) -> None:
        abnormal_dir = root / "Abnormal"
        if not abnormal_dir.is_dir():
            return
        abnormal = {path.name for path in abnormal_dir.glob("*.wav")}
        combined: set[str] = set()
        for label in ("Airway", "Lung_Parenchymal"):
            label_dir = root / label
            if not label_dir.is_dir():
                raise FileNotFoundError(
                    f"Required label directory not found: {label_dir}"
                )
            combined.update(path.name for path in label_dir.glob("*.wav"))
        if abnormal != combined:
            raise ValueError(
                "Abnormal overlay differs from Airway + Lung_Parenchymal. "
                f"root={root} extra={sorted(abnormal - combined)[:10]} "
                f"missing={sorted(combined - abnormal)[:10]}"
            )
        self.logger.info("Validated Abnormal overlay: %s", root)


class DiseaseDatasetCatalogBuilder:
    def __init__(
        self,
        *,
        data_root: Path,
        excel_root: Path,
        out_dir: Path,
    ):
        self.data_root = data_root
        self.excel_root = excel_root
        self.out_dir = out_dir

    def build(self) -> CatalogBuildResult:
        AbnormalOverlayValidator(self.data_root).validate()
        wav_inventory = WavInventoryLoader(self.data_root).load()
        annotations, warnings = ExcelAnnotationLoader(self.excel_root).load()
        memberships = SplitMembershipLoader(self.data_root).load(wav_inventory)
        rows: list[dict[str, Any]] = []
        used_annotations: set[str] = set()

        for stem in sorted(wav_inventory):
            wav = wav_inventory[stem]
            annotation = annotations.get(stem)
            membership = memberships.get(stem)
            notes: list[str] = []
            if annotation is None:
                notes.append("missing_excel_metadata")
                warnings.append(
                    CatalogWarning(
                        warning_type="wav_missing_excel_metadata",
                        filename_stem=stem,
                        filename=wav.filename,
                        primary_label=wav.primary_label,
                        path=str(wav.audio_path),
                        message="WAV file has no matching Excel metadata row",
                    )
                )
            else:
                used_annotations.add(stem)
                if annotation.source_label != wav.primary_label:
                    notes.append("excel_source_label_mismatch")
                    warnings.append(
                        CatalogWarning(
                            warning_type="primary_label_differs_from_excel_source",
                            filename_stem=stem,
                            filename=wav.filename,
                            primary_label=wav.primary_label,
                            source_label=annotation.source_label,
                            path=str(wav.audio_path),
                            excel_source_file=str(annotation.excel_source_file),
                            excel_row_number=annotation.excel_row_number,
                            message="Primary label differs from Excel source file label",
                        )
                    )
            if membership is None:
                notes.append("missing_split_membership")
                warnings.append(
                    CatalogWarning(
                        warning_type="wav_missing_split_membership",
                        filename_stem=stem,
                        filename=wav.filename,
                        primary_label=wav.primary_label,
                        path=str(wav.audio_path),
                        message="WAV file has no fold/test split membership",
                    )
                )
            elif membership.split_primary_label != wav.primary_label:
                notes.append("split_label_mismatch")
                warnings.append(
                    CatalogWarning(
                        warning_type="primary_label_differs_from_split_folder",
                        filename_stem=stem,
                        filename=wav.filename,
                        primary_label=wav.primary_label,
                        source_label=membership.split_primary_label,
                        path=str(membership.split_audio_path),
                        message="Primary label differs from split folder label",
                    )
                )
            rows.append(catalog_row(wav, annotation, membership, notes))

        for stem, annotation in sorted(annotations.items()):
            if stem in used_annotations:
                continue
            warnings.append(
                CatalogWarning(
                    warning_type="excel_row_missing_wav",
                    filename_stem=stem,
                    filename=annotation.filename,
                    source_label=annotation.source_label,
                    excel_source_file=str(annotation.excel_source_file),
                    excel_row_number=annotation.excel_row_number,
                    message="Excel row has no matching primary wav file",
                )
            )

        summary = build_summary(
            rows=rows,
            warnings=warnings,
            data_root=self.data_root,
            excel_root=self.excel_root,
            out_dir=self.out_dir,
        )
        return CatalogBuildResult(rows=rows, warnings=warnings, summary=summary)


class CatalogWriter(LoggingMixin):
    def __init__(self, out_dir: Path):
        self.out_dir = out_dir

    def write(self, result: CatalogBuildResult) -> dict[str, Path]:
        Fs.ensure_dir(self.out_dir)
        catalog_path = self.out_dir / "disease_dataset_catalog.csv"
        warnings_path = self.out_dir / "disease_dataset_catalog_warnings.csv"
        summary_path = self.out_dir / "disease_dataset_catalog_summary.json"
        write_csv(catalog_path, result.rows, CATALOG_COLUMNS)
        write_csv(
            warnings_path,
            [asdict(warning) for warning in result.warnings],
            WARNING_COLUMNS,
        )
        summary_path.write_text(
            json.dumps(result.summary, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        self.logger.info("Wrote %s", catalog_path)
        self.logger.info("Wrote %s", warnings_path)
        self.logger.info("Wrote %s", summary_path)
        return {
            "catalog": catalog_path,
            "warnings": warnings_path,
            "summary": summary_path,
        }


def catalog_row(
    wav: WavRecord,
    annotation: ExcelAnnotationRecord | None,
    membership: SplitMembership | None,
    notes: Sequence[str],
) -> dict[str, Any]:
    binary_label = "Normal" if wav.primary_label == "Normal" else "Abnormal"
    return {
        "filename_stem": wav.filename_stem,
        "filename": wav.filename,
        "audio_path": str(wav.audio_path),
        "primary_label": wav.primary_label,
        "binary_label": binary_label,
        "is_abnormal": binary_label == "Abnormal",
        "excel_source_file": ""
        if annotation is None
        else str(annotation.excel_source_file),
        "excel_row_number": "" if annotation is None else annotation.excel_row_number,
        "auscultation_raw": "" if annotation is None else annotation.auscultation.raw,
        "auscultation_code": ""
        if annotation is None
        else none_to_blank(annotation.auscultation.code),
        "auscultation_label": ""
        if annotation is None
        else annotation.auscultation.label,
        "diagnosis_raw": "" if annotation is None else annotation.diagnosis.raw,
        "diagnosis_code": ""
        if annotation is None
        else none_to_blank(annotation.diagnosis.code),
        "diagnosis_label": "" if annotation is None else annotation.diagnosis.label,
        "split": "" if membership is None else membership.split,
        "split_role": "" if membership is None else membership.split_role,
        "cv_validation_fold": ""
        if membership is None
        else none_to_blank(membership.cv_validation_fold),
        "is_test": "" if membership is None else membership.is_test,
        "split_primary_label": ""
        if membership is None
        else membership.split_primary_label,
        "split_audio_path": ""
        if membership is None
        else str(membership.split_audio_path),
        "appears_in_abnormal_overlay": ""
        if membership is None
        else membership.appears_in_abnormal_overlay,
        "abnormal_overlay_path": ""
        if membership is None or membership.abnormal_overlay_path is None
        else str(membership.abnormal_overlay_path),
        "has_excel_metadata": annotation is not None,
        "has_split_membership": membership is not None,
        "label_matches_excel_source": ""
        if annotation is None
        else wav.primary_label == annotation.source_label,
        "label_matches_split_folder": ""
        if membership is None
        else wav.primary_label == membership.split_primary_label,
        "notes": ";".join(notes),
    }


def build_summary(
    *,
    rows: Sequence[Mapping[str, Any]],
    warnings: Sequence[CatalogWarning],
    data_root: Path,
    excel_root: Path,
    out_dir: Path,
) -> dict[str, Any]:
    primary_counts = Counter(str(row["primary_label"]) for row in rows)
    binary_counts = Counter(str(row["binary_label"]) for row in rows)
    split_counts = Counter(str(row["split"]) for row in rows)
    auscultation_counts = Counter(
        str(row["auscultation_label"]) for row in rows if row["auscultation_label"]
    )
    diagnosis_counts = Counter(
        str(row["diagnosis_label"]) for row in rows if row["diagnosis_label"]
    )
    warning_counts = Counter(warning.warning_type for warning in warnings)
    return {
        "generated_at": datetime.now(UTC).isoformat(),
        "data_root": str(data_root),
        "excel_root": str(excel_root),
        "out_dir": str(out_dir),
        "row_count": len(rows),
        "primary_label_counts": dict(sorted(primary_counts.items())),
        "binary_label_counts": dict(sorted(binary_counts.items())),
        "split_counts": dict(sorted(split_counts.items())),
        "auscultation_label_counts": dict(sorted(auscultation_counts.items())),
        "diagnosis_label_counts": dict(sorted(diagnosis_counts.items())),
        "warning_counts": dict(sorted(warning_counts.items())),
        "warning_count": len(warnings),
    }


def write_csv(
    path: Path, rows: Sequence[Mapping[str, Any]], fieldnames: Sequence[str]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fieldnames))
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def normalize_filename_stem(value: object) -> str:
    raw = str(value).strip()
    stem = Path(raw).name
    if stem.lower().endswith(".wav"):
        stem = stem[:-4]
    return stem


def none_to_blank(value: object) -> object:
    return "" if value is None else value


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--excel-root", type=Path, default=DEFAULT_EXCEL_ROOT)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> None:
    args = parse_args(argv)
    Fs.ensure_dir(args.out_dir)
    enable_file_logging(args.out_dir / "disease_dataset_catalog_run.log", mode="w")
    result = DiseaseDatasetCatalogBuilder(
        data_root=args.data_root,
        excel_root=args.excel_root,
        out_dir=args.out_dir,
    ).build()
    paths = CatalogWriter(args.out_dir).write(result)
    logger.info(
        "Built disease dataset catalog | rows=%d | warnings=%d | catalog=%s",
        len(result.rows),
        len(result.warnings),
        paths["catalog"],
    )


if __name__ == "__main__":
    main()
